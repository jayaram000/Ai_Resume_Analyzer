import logging
from typing import Dict, Any
from docx import Document

logger = logging.getLogger(__name__)


class DOCXParsingService:
    """
    Parser for Microsoft Word (.docx) files using python-docx XML AST iteration.
    Extracts paragraphs and structured tables.
    """

    @staticmethod
    def extract_text(file_path: str) -> Dict[str, Any]:
        """
        Extracts structured text from a DOCX file.
        """
        raw_text_lines = []

        try:
            doc = Document(file_path)

            # 1. Iterate Paragraphs
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    raw_text_lines.append(text)

            # 2. Iterate Tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        raw_text_lines.append(" | ".join(row_text))

            full_text = "\n".join(raw_text_lines)

            return {
                "raw_text": full_text,
                "char_count": len(full_text),
                "is_scanned": False,
                "parser_used": "python-docx",
            }

        except Exception as e:
            logger.error(f"Error parsing DOCX file at '{file_path}': {str(e)}", exc_info=True)
            raise RuntimeError(f"Failed to parse DOCX document: {str(e)}") from e
