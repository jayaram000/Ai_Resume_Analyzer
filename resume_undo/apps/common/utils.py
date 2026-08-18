import os
import re
import zipfile
import io
from django.core.exceptions import ValidationError

def extract_text_from_pdf(file_obj):
    """
    Extracts text from a PDF file object.
    Tries PyMuPDF (fitz), pdfplumber, then pypdf, with ASCII regex fallback.
    """
    pdf_bytes = file_obj.read()
    bio = io.BytesIO(pdf_bytes)

    # 1. Try PyMuPDF (fitz)
    try:
        import fitz
        bio.seek(0)
        doc = fitz.open(stream=bio.read(), filetype="pdf")
        text = ""
        for page in doc:
            page_text = page.get_text()
            if page_text:
                text += page_text + "\n"
        if text.strip():
            return text
    except ImportError:
        pass
    except Exception:
        pass

    # 2. Try pdfplumber
    try:
        import pdfplumber
        bio.seek(0)
        text = ""
        with pdfplumber.open(bio) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        if text.strip():
            return text
    except ImportError:
        pass
    except Exception:
        pass

    # 3. Try pypdf
    try:
        import pypdf
        text = ""
        bio.seek(0)
        reader = pypdf.PdfReader(bio)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        if text.strip():
            return text
    except ImportError:
        pass
    except Exception:
        pass

    # 4. Fallback ASCII regex
    try:
        bio.seek(0)
        content = bio.read()
        ascii_strings = re.findall(b"[ -~]{4,}", content)
        return " ".join([s.decode('utf-8', errors='ignore') for s in ascii_strings])
    except Exception as e:
        raise ValidationError(f"Could not extract text from PDF: {str(e)}")

def extract_text_from_docx(file_obj):
    """
    Extracts text from a DOCX file object.
    Tries python-docx first, then falls back to direct XML parsing from zip.
    """
    docx_bytes = file_obj.read()
    bio = io.BytesIO(docx_bytes)

    try:
        import docx
        bio.seek(0)
        doc = docx.Document(bio)
        text = []
        for paragraph in doc.paragraphs:
            text.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text.append(cell.text)
        return "\n".join(text)
    except (ImportError, Exception):
        # Fallback: DOCX is a zip file containing XML.
        # We can extract word/document.xml and extract w:t tags.
        try:
            bio.seek(0)
            with zipfile.ZipFile(bio) as z:
                xml_content = z.read('word/document.xml').decode('utf-8')
                # Extract content from w:t tags
                text_runs = re.findall(r'<w:t.*?>(.*?)</w:t>', xml_content)
                # Decode basic XML entities
                text = " ".join(text_runs)
                text = text.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
                return text
        except Exception as e:
            raise ValidationError(f"Could not extract text from DOCX: {str(e)}")



class ExportService:
    @staticmethod
    def generate_pdf(report_type: str, data: dict, user_email: str) -> io.BytesIO:
        import io
        from django.utils import timezone
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        
        styles = getSampleStyleSheet()
        
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=22,
            leading=26,
            textColor=colors.HexColor('#6366F1'),
            spaceAfter=15
        )
        subtitle_style = ParagraphStyle(
            'DocSubtitle',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#94A3B8'),
            spaceAfter=25
        )
        h2_style = ParagraphStyle(
            'SectionHeader',
            parent=styles['Heading2'],
            fontSize=13,
            leading=17,
            textColor=colors.HexColor('#14B8A6'),
            spaceBefore=12,
            spaceAfter=8
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#1E293B')
        )
        bold_body_style = ParagraphStyle(
            'BoldBody',
            parent=body_style,
            fontName='Helvetica-Bold'
        )

        elements = []
        
        elements.append(Paragraph(f"AI CAREER COPILOT - {report_type.replace('_', ' ').upper()} REPORT", title_style))
        elements.append(Paragraph(f"Generated for: {user_email}  |  Date: {timezone.now().strftime('%Y-%m-%d')}", subtitle_style))
        elements.append(Spacer(1, 10))

        if report_type == "ats_report":
            elements.append(Paragraph("ATS Score Overview", h2_style))
            elements.append(Paragraph(f"ATS Score: <b>{data.get('ats_score', 0)} / 100</b>", body_style))
            elements.append(Spacer(1, 10))
            
            elements.append(Paragraph("Score Details", h2_style))
            table_data = [
                [Paragraph("<b>Category</b>", body_style), Paragraph("<b>Score</b>", body_style)],
                [Paragraph("Keyword Match", body_style), Paragraph(str(data.get('keyword_score', 0)), body_style)],
                [Paragraph("Formatting", body_style), Paragraph(str(data.get('formatting_score', 0)), body_style)],
                [Paragraph("Skills Match", body_style), Paragraph(str(data.get('skills_score', 0)), body_style)],
                [Paragraph("Experience Depth", body_style), Paragraph(str(data.get('experience_score', 0)), body_style)],
                [Paragraph("Education Context", body_style), Paragraph(str(data.get('education_score', 0)), body_style)],
                [Paragraph("Completeness", body_style), Paragraph(str(data.get('completeness_score', 0)), body_style)],
            ]
            t = Table(table_data, colWidths=[200, 100])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#1E293B')),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('BOTTOMPADDING', (0,0), (-1,0), 6),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')])
            ]))
            elements.append(t)
            elements.append(Spacer(1, 15))
            
            elements.append(Paragraph("Actionable Suggestions", h2_style))
            for sug in data.get("suggestions", []):
                elements.append(Paragraph(f"• {sug}", body_style))
                elements.append(Spacer(1, 4))
                
        elif report_type == "jd_match":
            elements.append(Paragraph("Job Description Match Score", h2_style))
            elements.append(Paragraph(f"Overall Match Score: <b>{data.get('match_score', 0)}%</b>", body_style))
            elements.append(Spacer(1, 10))
            
            elements.append(Paragraph("ATS Compatibility Status", h2_style))
            compat = data.get("ats_compatibility", {})
            elements.append(Paragraph(f"Status: <b>{compat.get('status', 'N/A')}</b> (ATS Compatibility Score: {compat.get('score', 0)})", body_style))
            elements.append(Spacer(1, 10))
            
            elements.append(Paragraph("Missing Keywords & Skills", h2_style))
            elements.append(Paragraph(f"<b>Missing Keywords:</b> {', '.join(data.get('missing_keywords', [])) or 'None'}", body_style))
            elements.append(Spacer(1, 6))
            elements.append(Paragraph(f"<b>Missing Skills:</b> {', '.join(data.get('missing_skills', [])) or 'None'}", body_style))
            elements.append(Spacer(1, 10))
            
            elements.append(Paragraph("Optimization Recommendations", h2_style))
            for rec in data.get("recommendations", []):
                elements.append(Paragraph(f"• {rec}", body_style))
                elements.append(Spacer(1, 4))

        elif report_type == "career_report":
            elements.append(Paragraph("Overall Career Readiness Score", h2_style))
            elements.append(Paragraph(f"Career Score: <b>{data.get('career_score', 0)} / 100</b>", body_style))
            elements.append(Spacer(1, 10))
            
            elements.append(Paragraph("Readiness Breakdown", h2_style))
            factors = data.get("factors") if isinstance(data.get("factors"), dict) else {}
            table_data = [
                [Paragraph("<b>Factor</b>", body_style), Paragraph("<b>Score</b>", body_style)],
                [Paragraph("ATS Resume Optimization (25% weight)", body_style), Paragraph(str(factors.get('ats', 0)), body_style)],
                [Paragraph("Skill Coverage vs Industry (20% weight)", body_style), Paragraph(str(factors.get('skills', 0)), body_style)],
                [Paragraph("Resume Completeness (15% weight)", body_style), Paragraph(str(factors.get('completeness', 0)), body_style)],
                [Paragraph("Experience Depth (15% weight)", body_style), Paragraph(str(factors.get('experience', 0)), body_style)],
                [Paragraph("Job Match Quality (15% weight)", body_style), Paragraph(str(factors.get('job_match', 0)), body_style)],
                [Paragraph("Profile Completeness (10% weight)", body_style), Paragraph(str(factors.get('profile', 0)), body_style)],
            ]
            t = Table(table_data, colWidths=[250, 80])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#1E293B')),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('BOTTOMPADDING', (0,0), (-1,0), 6),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')])
            ]))
            elements.append(t)
            elements.append(Spacer(1, 15))
            
            elements.append(Paragraph("Recommended Improvements", h2_style))
            import html
            for act in (data.get("recommended_actions") or []):
                elements.append(Paragraph(f"• {html.escape(str(act))}", body_style))
                elements.append(Spacer(1, 4))
                
        elif report_type == "skill_gap":
            elements.append(Paragraph(f"Skill Gap analysis for Target Role: {data.get('target_role', 'N/A')}", h2_style))
            elements.append(Paragraph(f"<b>Missing Skills identified:</b> {', '.join(data.get('missing_skills', [])) or 'None'}", body_style))
            elements.append(Spacer(1, 15))
            
            elements.append(Paragraph("Learning & Upskilling Priority Checklist", h2_style))
            table_data = [
                [Paragraph("<b>Skill</b>", body_style), Paragraph("<b>Priority</b>", body_style), Paragraph("<b>Time Estimate</b>", body_style), Paragraph("<b>Resources</b>", body_style)]
            ]
            for lp in data.get("learning_priority", []):
                table_data.append([
                    Paragraph(lp.get("skill", ""), body_style),
                    Paragraph(lp.get("priority", ""), bold_body_style),
                    Paragraph(lp.get("time_estimate", ""), body_style),
                    Paragraph(lp.get("resources", ""), body_style)
                ])
            t = Table(table_data, colWidths=[85, 65, 80, 200])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F1F5F9')),
                ('TEXTCOLOR', (0,0), (-1,0), colors.HexColor('#1E293B')),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('BOTTOMPADDING', (0,0), (-1,0), 6),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')])
            ]))
            elements.append(t)

        elif report_type == "roadmap":
            elements.append(Paragraph(f"Career Transition Roadmap: {data.get('current_role', 'N/A')} to {data.get('target_role', 'N/A')}", h2_style))
            elements.append(Spacer(1, 10))
            
            elements.append(Paragraph("Required Tech Stack", h2_style))
            elements.append(Paragraph(", ".join(data.get("technologies", [])), body_style))
            elements.append(Spacer(1, 10))
            
            elements.append(Paragraph("Recommended Premium Certifications", h2_style))
            elements.append(Paragraph(", ".join(data.get("certifications", [])), body_style))
            elements.append(Spacer(1, 10))
            
            elements.append(Paragraph("Phase-by-Phase Transition Plan", h2_style))
            for lp in data.get("learning_path", []):
                elements.append(Paragraph(f"<b>{lp.get('phase', '')}</b>", body_style))
                for ms in lp.get("milestones", []):
                    elements.append(Paragraph(f"• {ms}", body_style))
                    elements.append(Spacer(1, 2))
                elements.append(Spacer(1, 8))
                
        else:
            elements.append(Paragraph("Report Details", h2_style))
            elements.append(Paragraph(str(data), body_style))

        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_docx(report_type: str, data: dict, user_email: str) -> io.BytesIO:
        import io
        from django.utils import timezone
        import docx
        
        doc = docx.Document()
        doc.add_heading(f"AI Career Copilot - {report_type.replace('_', ' ').upper()} Report", 0)
        doc.add_paragraph(f"Generated for: {user_email}  |  Date: {timezone.now().strftime('%Y-%m-%d')}")

        if report_type == "ats_report":
            doc.add_heading(f"ATS Score: {data.get('ats_score', 0)} / 100", level=1)
            
            table = doc.add_table(rows=1, cols=2)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Category'
            hdr_cells[1].text = 'Score'
            
            breakdown = [
                ("Keyword Match", data.get('keyword_score', 0)),
                ("Formatting", data.get('formatting_score', 0)),
                ("Skills Match", data.get('skills_score', 0)),
                ("Experience Depth", data.get('experience_score', 0)),
                ("Education Context", data.get('education_score', 0)),
                ("Completeness", data.get('completeness_score', 0)),
            ]
            for name, score in breakdown:
                row_cells = table.add_row().cells
                row_cells[0].text = name
                row_cells[1].text = str(score)
                
            doc.add_heading("Suggestions for Improvement", level=1)
            for sug in data.get("suggestions", []):
                doc.add_paragraph(sug, style='List Bullet')
                
        elif report_type == "jd_match":
            doc.add_heading(f"Overall Match Score: {data.get('match_score', 0)}%", level=1)
            doc.add_paragraph(f"ATS Status: {data.get('ats_compatibility', {}).get('status', 'N/A')}")
            
            doc.add_heading("Skill & Keyword Gaps", level=1)
            doc.add_paragraph(f"Missing Keywords: {', '.join(data.get('missing_keywords', [])) or 'None'}")
            doc.add_paragraph(f"Missing Skills: {', '.join(data.get('missing_skills', [])) or 'None'}")
            
            doc.add_heading("Recommendations", level=1)
            for rec in data.get("recommendations", []):
                doc.add_paragraph(rec, style='List Bullet')

        elif report_type == "career_report":
            doc.add_heading(f"Career Score: {data.get('career_score', 0)} / 100", level=1)
            
            table = doc.add_table(rows=1, cols=2)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Readiness Factor'
            hdr_cells[1].text = 'Score'
            
            factors = data.get("factors", {})
            rows = [
                ("ATS Resume Score", factors.get('ats', 0)),
                ("Skill Coverage", factors.get('skills', 0)),
                ("Resume Completeness", factors.get('completeness', 0)),
                ("Experience Strength", factors.get('experience', 0)),
                ("Job Match Quality", factors.get('job_match', 0)),
                ("Profile Completeness", factors.get('profile', 0)),
            ]
            for name, val in rows:
                row_cells = table.add_row().cells
                row_cells[0].text = name
                row_cells[1].text = str(val)

            doc.add_heading("Recommended Goals", level=1)
            for act in data.get("recommended_actions", []):
                doc.add_paragraph(act, style='List Bullet')
                
        elif report_type == "skill_gap":
            doc.add_heading(f"Upskilling Action Plan: {data.get('target_role', 'N/A')}", level=1)
            doc.add_paragraph(f"Missing Skills: {', '.join(data.get('missing_skills', [])) or 'None'}")
            
            table = doc.add_table(rows=1, cols=4)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Skill'
            hdr_cells[1].text = 'Priority'
            hdr_cells[2].text = 'Time Estimate'
            hdr_cells[3].text = 'Learning Resource'
            
            for lp in data.get("learning_priority", []):
                row_cells = table.add_row().cells
                row_cells[0].text = lp.get("skill", "")
                row_cells[1].text = lp.get("priority", "")
                row_cells[2].text = lp.get("time_estimate", "")
                row_cells[3].text = lp.get("resources", "")

        elif report_type == "roadmap":
            doc.add_heading(f"Transition Roadmap: {data.get('current_role', 'N/A')} to {data.get('target_role', 'N/A')}", level=1)
            doc.add_paragraph(f"Target Technologies: {', '.join(data.get('technologies', []))}")
            doc.add_paragraph(f"Certifications: {', '.join(data.get('certifications', []))}")
            
            doc.add_heading("Step-by-Step Milestones", level=1)
            for lp in data.get("learning_path", []):
                doc.add_heading(lp.get("phase", ""), level=2)
                for ms in lp.get("milestones", []):
                    doc.add_paragraph(ms, style='List Bullet')

        else:
            doc.add_paragraph(str(data))

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

